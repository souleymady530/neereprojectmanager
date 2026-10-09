"""
Utilitaires pour l'import de tickets depuis un fichier Excel (.xlsx).

Colonnes reconnues (insensibles à la casse, espaces ignorés) :
  Sujet*        – texte libre (obligatoire)
  Description   – texte long
  Tracker       – nom du tracker (ex. Bug, Fonctionnalité)
  Priorité      – nom de la priorité (ex. Haute, Normale)
  Statut        – nom du statut
  Assigné à     – username ou email
  Date début    – YYYY-MM-DD ou DD/MM/YYYY
  Date échéance – YYYY-MM-DD ou DD/MM/YYYY
  Heures estimées – nombre décimal
  Avancement    – entier 0–100
  Version cible – nom de la version du projet
"""

from __future__ import annotations

import datetime
import io
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


# ── Colonnes reconnues ────────────────────────────────────────────────────────
COLUMN_ALIASES: dict[str, str] = {
    # clé normalisée : nom interne
    "sujet":            "subject",
    "subject":          "subject",
    "titre":            "subject",
    "title":            "subject",
    "description":      "description",
    "tracker":          "tracker",
    "priorité":         "priority",
    "priorite":         "priority",
    "priority":         "priority",
    "statut":           "status",
    "status":           "status",
    "assignéà":         "assigned_to",
    "assignea":         "assigned_to",
    "assigné":          "assigned_to",
    "assignedto":       "assigned_to",
    "responsable":      "assigned_to",
    "datedébut":        "start_date",
    "datedebut":        "start_date",
    "startdate":        "start_date",
    "dateéchéance":     "due_date",
    "dateecheance":     "due_date",
    "duedate":          "due_date",
    "échéance":         "due_date",
    "echeance":         "due_date",
    "heuresestimées":   "estimated_hours",
    "heuresestimees":   "estimated_hours",
    "estimatedhours":   "estimated_hours",
    "heures":           "estimated_hours",
    "avancement":       "done_ratio",
    "avancement(%)":    "done_ratio",
    "doneratio":        "done_ratio",
    "versioncible":     "fixed_version",
    "version":          "fixed_version",
    "fixedversion":     "fixed_version",
}


def _normalize(header: str) -> str:
    """Normalise un en-tête : minuscules, sans espaces ni caractères spéciaux."""
    return "".join(c for c in header.lower() if c.isalnum() or c in "éèêëàâùûîïôœç").replace(" ", "")


@dataclass
class ImportRow:
    row_num: int
    data: dict[str, Any]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def is_valid(self):
        return not self.errors


@dataclass
class ImportResult:
    rows: list[ImportRow] = field(default_factory=list)
    created: int = 0
    skipped: int = 0
    parse_error: str | None = None


def _parse_date(value) -> datetime.date | None:
    if value is None:
        return None
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.date() if isinstance(value, datetime.datetime) else value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_excel(file_obj) -> ImportResult:
    """
    Lit le fichier Excel et retourne un ImportResult avec les lignes parsées.
    Ne touche pas à la base de données.
    """
    result = ImportResult()
    try:
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as e:
        result.parse_error = f"Impossible de lire le fichier : {e}"
        return result

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        result.parse_error = "Le fichier est vide."
        return result

    # Détecter les en-têtes (première ligne non vide)
    header_row = None
    data_start = 0
    for i, row in enumerate(rows):
        if any(cell is not None for cell in row):
            header_row = [str(c).strip() if c is not None else "" for c in row]
            data_start = i + 1
            break

    if header_row is None:
        result.parse_error = "Aucun en-tête trouvé."
        return result

    # Mapper les colonnes
    col_map: dict[int, str] = {}
    for idx, h in enumerate(header_row):
        key = _normalize(h)
        internal = COLUMN_ALIASES.get(key)
        if internal:
            col_map[idx] = internal

    if "subject" not in col_map.values():
        result.parse_error = (
            "Colonne « Sujet » introuvable. "
            "Vérifiez que votre fichier contient bien une colonne nommée « Sujet »."
        )
        return result

    for row_idx, row in enumerate(rows[data_start:], start=data_start + 2):
        # Ignorer les lignes totalement vides
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw: dict[str, Any] = {}
        for col_idx, internal_name in col_map.items():
            val = row[col_idx] if col_idx < len(row) else None
            raw[internal_name] = val

        import_row = ImportRow(row_num=row_idx, data={})
        errors = import_row.errors
        warnings = import_row.warnings

        # subject (obligatoire)
        subject = str(raw.get("subject") or "").strip()
        if not subject:
            errors.append("La colonne « Sujet » est vide.")
        else:
            import_row.data["subject"] = subject[:255]

        # description
        desc = raw.get("description")
        import_row.data["description"] = str(desc).strip() if desc is not None else ""

        # tracker (résolu plus tard dans save_rows)
        import_row.data["_tracker_name"] = str(raw.get("tracker") or "").strip()

        # priority
        import_row.data["_priority_name"] = str(raw.get("priority") or "").strip()

        # status
        import_row.data["_status_name"] = str(raw.get("status") or "").strip()

        # assigned_to
        import_row.data["_assigned_to"] = str(raw.get("assigned_to") or "").strip()

        # dates
        start = _parse_date(raw.get("start_date"))
        due = _parse_date(raw.get("due_date"))
        if raw.get("start_date") and start is None:
            warnings.append(f"Date début non reconnue : {raw['start_date']!r}")
        if raw.get("due_date") and due is None:
            warnings.append(f"Date échéance non reconnue : {raw['due_date']!r}")
        import_row.data["start_date"] = start
        import_row.data["due_date"] = due

        # estimated_hours
        eh = raw.get("estimated_hours")
        if eh is not None and str(eh).strip():
            try:
                import_row.data["estimated_hours"] = round(float(eh), 2)
            except (ValueError, TypeError):
                warnings.append(f"Heures estimées ignorées : {eh!r}")
        else:
            import_row.data["estimated_hours"] = None

        # done_ratio
        dr = raw.get("done_ratio")
        if dr is not None and str(dr).strip():
            try:
                ratio = int(float(dr))
                import_row.data["done_ratio"] = max(0, min(100, ratio))
            except (ValueError, TypeError):
                warnings.append(f"Avancement ignoré : {dr!r}")
        else:
            import_row.data["done_ratio"] = 0

        # fixed_version
        import_row.data["_version_name"] = str(raw.get("fixed_version") or "").strip()

        result.rows.append(import_row)

    return result


def save_rows(result: ImportResult, project, author) -> tuple[int, int, list[str]]:
    """
    Crée les tickets en base à partir des lignes parsées valides.
    Retourne (created_count, skipped_count, global_warnings).
    """
    from .models import Issue, Tracker, IssueStatus, IssuePriority
    from django.contrib.auth import get_user_model
    from versions.models import Version

    User = get_user_model()

    # Caches pour éviter les requêtes répétées
    trackers   = {t.name.lower(): t for t in Tracker.objects.all()}
    statuses   = {s.name.lower(): s for s in IssueStatus.objects.all()}
    priorities = {p.name.lower(): p for p in IssuePriority.objects.all()}
    versions   = {v.name.lower(): v for v in Version.objects.filter(project=project)}
    users      = {u.username.lower(): u for u in User.objects.all()}
    users.update({u.email.lower(): u for u in User.objects.all() if u.email})

    default_tracker  = Tracker.objects.order_by("position").first()
    default_priority = IssuePriority.objects.filter(is_default=True).first() \
                       or IssuePriority.objects.order_by("position").first()
    default_status   = (
        default_tracker.default_status if default_tracker and default_tracker.default_status
        else IssueStatus.objects.order_by("position").first()
    )

    created = 0
    skipped = 0
    global_warnings: list[str] = []

    for row in result.rows:
        if not row.is_valid:
            skipped += 1
            continue

        d = row.data

        # Résoudre les références par nom
        tracker_name = d.get("_tracker_name", "").lower()
        tracker = trackers.get(tracker_name) or default_tracker
        if tracker_name and not trackers.get(tracker_name):
            global_warnings.append(f"Ligne {row.row_num} : tracker « {d['_tracker_name']} » inconnu, utilisation du tracker par défaut.")

        priority_name = d.get("_priority_name", "").lower()
        priority = priorities.get(priority_name) or default_priority
        if priority_name and not priorities.get(priority_name):
            global_warnings.append(f"Ligne {row.row_num} : priorité « {d['_priority_name']} » inconnue, utilisation de la priorité par défaut.")

        status_name = d.get("_status_name", "").lower()
        status = statuses.get(status_name)
        if not status:
            status = (tracker.default_status if tracker and tracker.default_status else default_status)
        if status_name and not statuses.get(status_name):
            global_warnings.append(f"Ligne {row.row_num} : statut « {d['_status_name']} » inconnu, utilisation du statut par défaut.")

        assigned_raw = d.get("_assigned_to", "").lower()
        assigned_to = users.get(assigned_raw) if assigned_raw else None
        if assigned_raw and not assigned_to:
            global_warnings.append(f"Ligne {row.row_num} : utilisateur « {d['_assigned_to']} » introuvable, ticket non assigné.")

        version_name = d.get("_version_name", "").lower()
        fixed_version = versions.get(version_name) if version_name else None
        if version_name and not fixed_version:
            global_warnings.append(f"Ligne {row.row_num} : version « {d['_version_name']} » introuvable, ignorée.")

        if not tracker or not status or not priority:
            skipped += 1
            global_warnings.append(f"Ligne {row.row_num} : données de référence manquantes, ligne ignorée.")
            continue

        Issue.objects.create(
            project=project,
            author=author,
            subject=d["subject"],
            description=d.get("description", ""),
            tracker=tracker,
            status=status,
            priority=priority,
            assigned_to=assigned_to,
            fixed_version=fixed_version,
            start_date=d.get("start_date"),
            due_date=d.get("due_date"),
            estimated_hours=d.get("estimated_hours"),
            done_ratio=d.get("done_ratio", 0),
        )
        created += 1

    result.created = created
    result.skipped = skipped
    return created, skipped, global_warnings


def build_template_xlsx() -> bytes:
    """Retourne un fichier .xlsx de modèle vide avec les colonnes attendues."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets"

    headers = [
        "Sujet",
        "Description",
        "Tracker",
        "Priorité",
        "Statut",
        "Assigné à",
        "Date début",
        "Date échéance",
        "Heures estimées",
        "Avancement",
        "Version cible",
    ]

    # Style en-tête
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="3D4B5A")
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    # Ligne d'exemple
    ws.append([
        "Exemple de ticket",
        "Description détaillée ici",
        "Bug",
        "Normale",
        "Nouveau",
        "admin",
        "2026-10-15",
        "2026-10-30",
        "4.5",
        "0",
        "",
    ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
